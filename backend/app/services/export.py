import csv
import io

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.schemas import PermissionsMatrix


def matrix_to_csv(matrix: PermissionsMatrix) -> bytes:
    buffer = io.StringIO()
    writer = csv.writer(buffer)
    writer.writerow(["Ruta", *[role.name for role in matrix.roles]])
    for folder in matrix.folders:
        writer.writerow([folder.path, *[folder.permissions.get(role.id, "") for role in matrix.roles]])
    return buffer.getvalue().encode("utf-8-sig")


def matrix_to_xlsx(matrix: PermissionsMatrix) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Permisos"
    headers = ["Ruta", *[role.name for role in matrix.roles]]
    sheet.append(headers)
    for folder in matrix.folders:
        sheet.append([folder.path, *[folder.permissions.get(role.id, "") for role in matrix.roles]])

    header_fill = PatternFill("solid", fgColor="1F6B4A")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    for index, _header in enumerate(headers, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = 36 if index == 1 else 16
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{sheet.max_row}"

    output = io.BytesIO()
    workbook.save(output)
    return output.getvalue()
